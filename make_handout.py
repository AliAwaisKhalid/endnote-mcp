"""Build EndNote_MCP_Handout.docx (A4, two pages), in the style of the Zotero handout."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"J:\Claude\endnote-mcp\EndNote_MCP_Handout.docx"
NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x59, 0x59, 0x59)

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.9)
sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.5)
st = d.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(3)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.08


def shade(pr, fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    pr.append(sh)


def para(text="", bold=False, size=None, color=None, italic=False, after=None, before=None):
    p = d.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    return p


def heading(text):
    p = para(text, bold=True, size=12, color=NAVY, before=9, after=3)
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "8EA9DB")):
        b.set(qn("w:" + k), v)
    bd.append(b)
    pPr.append(bd)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(parts):
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.left_indent = Cm(0.55)
    for t, b in parts if isinstance(parts, list) else [(parts, False)]:
        r = p.add_run(t)
        r.bold = b
    return p


def code(lines):
    for ln in lines:
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(0)
        shade(p._p.get_or_add_pPr(), "F2F2F2")
        r = p.add_run(ln)
        r.font.name = "Consolas"
        r.font.size = Pt(9)
        r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
    d.add_paragraph().paragraph_format.space_after = Pt(1)


def table(header, rows, widths):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.width = Cm(widths[i])
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c._tc.get_or_add_tcPr(), "1F3864")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].width = Cm(widths[i])
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(v)
            r.font.size = Pt(8.5)
            if i == 0:
                r.bold = True
                r.font.name = "Consolas"
    d.add_paragraph().paragraph_format.space_after = Pt(1)


# ---- page 1
para("EndNote MCP: read-only access to EndNote 2025 libraries", bold=True, size=17, color=NAVY, after=2)
para("Claude Code can search, read and check your EndNote libraries. It never changes a library, a .Data folder or an attachment.",
     color=GREY, italic=True, after=6)

heading("Set up")
bullet([("Registered already. ", True), ("The server is in your user config as ", False), ("endnote", True),
        (". Check it with ", False), ("claude mcp list", True), (" (it should show Connected).", False)])
bullet([("Libraries are listed in ", False), ("libraries.json", True),
        (" next to server.py. Add a library there with its .enl path, then start a new Claude Code session so the list refreshes.", False)])
bullet([("EndNote can stay open. ", True),
        ("The server copies a library only after EndNote has saved it. If a copy fails, the server shows the last copy with its time.", False)])

heading("Tools")
table(["Tool", "What it does"], [
    ("list_libraries", "Configured libraries, reference counts, snapshot time."),
    ("search_items", "Words in author, title, year, keywords, abstract. Filters: group, keyword, type. fulltext=True searches PDF text."),
    ("find_by_identifier", "A DOI or doi.org link, a key such as Sports:12, a record number, or title words (small spelling mistakes allowed)."),
    ("get_item", "Every non-empty field, PDF paths (with whether each exists), groups and colour tags."),
    ("get_notes", "Notes and research notes for one record, or search them by words."),
    ("get_annotations", "Highlights and comments in attached PDFs, with page, text, colour and author."),
    ("get_fulltext", "Text of the first attached PDF, 30,000 characters by default. Cached."),
    ("list_collections", "Groups with type and item count. Smart groups show their cached members."),
    ("collection_items", "Members of one group, by name or group id."),
    ("list_tags", "Keywords with counts, or EndNote colour tags with counts (kind=colour)."),
    ("recent_items", "Newest records by date added or last updated (UTC)."),
    ("find_duplicates", "Records sharing a DOI, or the same title and year."),
    ("list_by_status", "Records whose keywords or Label start with a prefix such as replication or verify."),
    ("method_matrix", "Method, Estimator, Sample and Data lines from research notes or custom fields."),
    ("get_bibliography", "Formatted bibliography with pandoc and a CSL style such as APA."),
    ("export_items", "RIS, BibTeX or CSL-JSON for keys or a group. The file is saved and its text returned."),
    ("preview_import", "Checks an RIS file against the library. Writes only new records to a fresh RIS file for File > Import."),
    ("check_records", "Flags records that would cite badly in Word: split organisation names, missing year, author, title or journal, bad or shared DOI, reversed pages, all-capital titles."),
    ("check_docx_citations", "Reads the EndNote citations in a Word .docx (Word stays closed) and checks them against a library: missing or trashed records, author or year that differ, wrong or mixed libraries."),
    ("missing_pdfs", "Records with no PDF attached, a PDF missing on disk, or an attachment path outside the library."),
    ("export_annotations", "Highlights and comments of several records, saved as one Markdown file."),
    ("compare_libraries", "Records in one library that are not in another, by DOI or title and year."),
    ("watch_library", "Records added, updated, trashed, restored or removed since your last call."),
    ("get_online_record", "Full metadata for one DOI from Crossref or OpenAlex, and whether you already have it."),
    ("search_online", "Searches Crossref or OpenAlex by words; each result is marked IN LIBRARY or NEW."),
    ("online_citations", "Works a paper cites, or works citing it (OpenAlex), marked IN LIBRARY or NEW."),
    ("import_online", "DOIs to a RIS file of the new ones, after a duplicate check. You import it with File > Import."),
], widths=[3.6, 13.6])

# ---- page 2
d.add_page_break()
heading("Common requests")
bullet([("Is this paper already in my library? ", True), ("Ask by DOI or title. ", False),
        ("find_by_identifier", True), (" returns exact matches first.", False)])
bullet([("What did I highlight in this paper? ", True), ("Give its key. ", False), ("get_annotations", True),
        (" reads the PDF in place, and nothing is saved.", False)])
bullet([("Check an import before I make it. ", True), ("Give the RIS file. ", False), ("preview_import", True),
        (" marks each record NEW or DUPLICATE and writes the new ones to one file. You run the import in EndNote with File > Import.", False)])
bullet([("Check records before citing. ", True), ("Give keys or a group to ", False), ("check_records", True),
        (". An organisation written as several author entries (the Traveling Library problem) is flagged, so you fix it in EndNote before the Word update.", False)])
bullet([("Bring in a paper's references or citing papers. ", True), ("Give its DOI to ", False), ("online_citations", True),
        (", then the DOIs marked NEW to ", False), ("import_online", True), (". Import the RIS file in EndNote.", False)])
bullet([("Check a Word document's citations. ", True), ("Give the .docx path and the library to ", False), ("check_docx_citations", True),
        (". It lists citations whose record is missing, in the trash, or has a different author or year. Word does not need to be open.", False)])
bullet([("Build a methods table. ", True), ("Put Method:, Estimator:, Sample: and Data: lines in the research notes, then run ", False),
        ("method_matrix", True), (" for a group.", False)])
bullet([("Format references. ", True), ("Give keys to ", False), ("get_bibliography", True),
        (" (for example style apa), or a group to ", False), ("export_items", True), (".", False)])

heading("Keys and libraries")
para("A key is the library name, a colon and the record number, for example Sports:12. Record numbers are unique only within one "
     "library, so the library name is always part of the key. Every tool also takes a library argument, which may be a name or a path. "
     "The first library in libraries.json is the default.", after=4)
code(["Sports:12                              record 12 in the Sports library",
      "find_by_identifier  identifier=\"10.1186/s13102-024-00914-5\"  library=\"Sports\""])

heading("Limits")
bullet("Read only. Adding or changing records and tags needs EndNote itself.")
bullet("Results come from the last copy. A change in EndNote shows on the next call after EndNote saves it.")
bullet("A library with no Data\\sdb\\sdb.eni file is an older format. Convert it in EndNote 2025 first: File > Open > Open Library.")
bullet(".enlx files are not read. Save or export the library as .enl first.")
bullet("Smart groups show the members EndNote last cached, which may be out of date. Group sets were not found in the test libraries.")
bullet("A search that finds nothing returns one row with a message, not an error. search_items needs exact words; find_by_identifier tolerates small spelling mistakes. limit must be 1 or more.")
bullet("Attachment paths that lead outside the library's PDF folder are refused. preview_import reads only .ris or .txt files up to 10 MB. check_docx_citations reads only .docx files up to 100 MB.")
bullet("Sort order for text is an approximation of EndNote's own. Matching by words and DOIs is exact.")
bullet("PDF annotations are read with PyMuPDF. Highlighted text may differ slightly from EndNote's display.")

heading("Tested and reviewed (28 October 2026)")
bullet([("Tests: ", True), ("203 pytest tests, and test_mcp_client.py calls all 27 tools over MCP (52 of 52 pass offline, 61 of 61 with the online tests; library files unchanged).", False)])
bullet([("User simulations: ", True), ("31 common requests run end to end. Seven unclear messages were fixed (SIMULATION_FINDINGS.md).", False)])
bullet([("Security review: ", True), ("two fixes (PDF paths leaving the library, import file type and size) and two accepted risks (style path, error text). See SECURITY_FINDINGS.md.", False)])
bullet([("Word proposal v6: ", True), ("106 citations to 36 records, all match the Commodities library.", False)])
bullet([("Next: ", True), ("see NEXT_SESSION.md.", False)])

heading("Where things are")
bullet([("Copies of libraries: ", True), ("%LOCALAPPDATA%\\endnote-mcp\\cache\\", False)])
bullet([("Exports and RIS files: ", True), ("%LOCALAPPDATA%\\endnote-mcp\\work\\", False)])
bullet([("Server, tests and notes: ", True), ("J:\\Claude\\endnote-mcp\\ (server.py, test files, README.md, SCHEMA.md, SECURITY_FINDINGS.md)", False)])
bullet([("Privacy: ", True), ("Everything runs on this computer except the four online tools, which send only a DOI or search words to Crossref or OpenAlex. No library content leaves. Set ENDNOTE_MCP_OFFLINE=1 to turn them off. Google Scholar and Scopus are not used (no official API; no key).", False)])

d.save(OUT)
print("saved", OUT)
